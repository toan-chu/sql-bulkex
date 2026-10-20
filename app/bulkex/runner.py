"""One runner pass: rescan tables, refresh the page's field list, run every new request in input/.

Called every 2 minutes by Windows Task Scheduler (`pythonw app\\bulkex once`).
"""
from __future__ import annotations

import csv
import heapq
import datetime as dt
import json
import logging
import os
import re
import shutil
import time
import traceback
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from . import query as Q
from .catalog import build_catalog, load_defs, write_js
from .config import connect, load_yaml
from .onedrive import free_up, keep_workspace, online_only
from .scan import field_types, scan

log = logging.getLogger("bulkex")
STATUS_KEEP = 300  # newest finished requests shown on the page
LOCK_STALE_SECONDS = 3 * 3600


def setup_logging(log_dir: Path):
    log_dir.mkdir(parents=True, exist_ok=True)
    if not any(isinstance(h, logging.FileHandler) and Path(h.baseFilename) == (log_dir / "bulkex.log").resolve()
               for h in log.handlers):
        fh = logging.FileHandler(log_dir / "bulkex.log", encoding="utf-8")
        fh.setFormatter(logging.Formatter("[%(asctime)s] %(levelname)s %(message)s", "%Y-%m-%d %H:%M:%S"))
        log.addHandler(fh)
    log.setLevel(logging.INFO)


def slug(s, n=40):
    s = re.sub(r"[^\w\-]+", "-", str(s or ""), flags=re.UNICODE).strip("-")
    return s[:n] or "request"


# ---------------------------------------------------------------- status file (read by the page)
class DatabaseDown(Exception):
    """No data table could be reached: keep the old field list and leave requests waiting."""


class Status:
    """What the page shows under Recent requests.

    Finished requests are read straight from _system/done and _system/failed (each file carries its own
    result under "_result"), so the admin clears the list for everyone by deleting files there.
    Only requests still running or waiting are remembered in status.json.
    """

    def __init__(self, system_dir: Path):
        self.system_dir = system_dir
        self.input_dir = system_dir.parent / "input"
        self.json_path = system_dir / "status.json"
        self.js_path = system_dir / "status.js"
        try:
            data = json.loads(self.json_path.read_text(encoding="utf-8")) if self.json_path.exists() else {}
        except (ValueError, OSError):
            data = {}
        self.active = [r for r in (data.get("active") or []) if r.get("state") in ("running", "waiting")]
        self.items = []

    def put(self, rec: dict):
        self.active = [r for r in self.active if r.get("id") != rec.get("id")]
        if rec.get("state") in ("running", "waiting"):
            self.active.insert(0, dict(rec))
        self.save()

    def finished(self):
        out = []
        for folder in ("done", "failed"):
            for f in (self.system_dir / folder).glob("*.json"):
                try:
                    rec = json.loads(f.read_text(encoding="utf-8-sig")).get("_result")
                except (ValueError, OSError, AttributeError):
                    rec = None
                if not isinstance(rec, dict):  # a file from before v7.1, or one put there by hand
                    rec = {"id": f.stem, "name": f.stem, "message": "",
                           "finished_at": dt.datetime.fromtimestamp(f.stat().st_mtime).isoformat(timespec="seconds")}
                rec["state"] = folder
                out.append((rec.get("finished_at") or "", f.stat().st_mtime_ns, rec))
        out.sort(key=lambda x: x[:2], reverse=True)
        return [r for _, _, r in out[:STATUS_KEEP]]

    def save(self):
        done = self.finished()
        ids = {r.get("id") for r in done}
        # forget running/waiting records that have finished, or whose file left input/ (deleted by hand)
        self.active = [r for r in self.active if r.get("id") not in ids
                       and (not r.get("_file") or (self.input_dir / r["_file"]).exists())]
        self.items = self.active + done
        now = dt.datetime.now().isoformat(timespec="seconds")
        tmp = self.json_path.with_suffix(".json.tmp")
        tmp.write_text(json.dumps({"schema": "bulkex-status/2", "updated_at": now, "active": self.active},
                                  ensure_ascii=False, indent=1), encoding="utf-8")
        tmp.replace(self.json_path)
        write_js(self.js_path, "BULKEX_STATUS", {"schema": "bulkex-status/2", "updated_at": now, "requests": self.items},
                 "Written by the BulkEx runner from _system/done and _system/failed.")


# ---------------------------------------------------------------- Excel writing
BOLD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="4D148C")


def _cell_value(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, dt.datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    return v


def _row(ws, values):
    out = []
    for v in values:
        v = _cell_value(v)
        if isinstance(v, str) and v.startswith("="):  # keep text that looks like a formula as plain text
            c = WriteOnlyCell(ws, value=v)
            c.data_type = "s"
            out.append(c)
        else:
            out.append(v)
    return out


def _header(ws, names):
    cells = []
    for n in names:
        c = WriteOnlyCell(ws, value=n)
        c.font, c.fill = BOLD, HEAD_FILL
        cells.append(c)
    ws.append(cells)


def _write_csv(path: Path, heads, rows):
    """UTF-8 with BOM so Excel opens Vietnamese text correctly."""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(heads)
        for r in rows:
            w.writerow([_cell_value(v) for v in r])
    return path


# ---------------------------------------------------------------- the runner
class Runner:
    def __init__(self, settings, cfg, config_dir: Path):
        self.s = settings
        self.cfg = cfg
        self.config_dir = Path(config_dir)
        self.labels, self.defs = load_defs(settings.app_dir)
        self.found, self.catalog, self.types = {}, {}, {}

    # -- catalog -------------------------------------------------------------------------------
    def refresh_catalog(self):
        found, warnings = scan(self.cfg, self.defs.get("datasets") or {}, self.s.databases or None)
        if not any(found.values()):
            # never replace the page's field list with an empty one: the database is down or unreachable
            for w in warnings:
                log.warning(w)
            raise DatabaseDown("; ".join(warnings) or "no data tables found")
        self.found = found
        values = load_yaml(self.config_dir / "values.json", {})
        self.catalog = build_catalog(self.found, self.labels, self.defs, values)
        nat = self.defs.get("numbers_stored_as_text") or []
        self.types = {ds: field_types(t, nat) for ds, t in self.found.items() if t}
        write_js(self.s.system_dir / "catalog.js", "BULKEX_CATALOG", self.catalog, "Written by the BulkEx runner. Do not edit.")
        snap = {ds: [{k: t[k] for k in ("db", "schema", "table", "period")} for t in ts] for ds, ts in self.found.items()}
        try:
            old = json.loads((self.s.system_dir / "tables.json").read_text(encoding="utf-8")).get("warnings")
        except (OSError, ValueError, AttributeError):
            old = None
        (self.s.system_dir / "tables.json").write_text(json.dumps({"tables": snap, "warnings": warnings}, ensure_ascii=False, indent=1),
                                                        encoding="utf-8")
        if warnings != old:  # the same notes every 2 minutes would only bury real errors
            for w in warnings:
                log.warning(w)
        return warnings

    def label(self, ds, col):
        for f in self.catalog["datasets"].get(ds, {}).get("fields", []):
            if f["key"] == col:
                return f["label"]
        return col

    # -- one request ---------------------------------------------------------------------------
    def validate(self, req):
        if not isinstance(req, dict) or req.get("schema") != "bulkex-request/1":
            raise Q.RequestError("This file is not a BulkEx request. Send it again from the page.")
        if req.get("action") not in ("count", "export"):
            raise Q.RequestError("Unknown action in the request.")
        parts = req.get("parts")
        if not str(req.get("name") or "").strip() or not str(req.get("requester") or "").strip():
            raise Q.RequestError("Add a request name and your name, then send it again.")
        if not isinstance(parts, list) or not 0 < len(parts) <= 2:
            raise Q.RequestError("The request has no data set.")
        seen = set()
        for p in parts:
            ds = p.get("dataset")
            if ds not in self.types or ds in seen:
                raise Q.RequestError(f"Data set '{ds}' is not available on this machine.")
            seen.add(ds)
            labels = {f["key"]: f["label"] for f in self.catalog["datasets"][ds]["fields"]}
            Q.validate_part(p, {"types": self.types[ds]}, labels)
            rows = (p.get("output") or {}).get("rows")
            if rows:  # a filtered or sorted field always ends up in the file
                cols = list(dict.fromkeys(rows["columns"]))
                for k in Q.filter_fields(p["filter"]):
                    if k not in cols:
                        cols.append(k)
                for s in rows.get("sort") or []:
                    if s["field"] not in cols:
                        cols.append(s["field"])
                rows["columns"] = cols
        return req

    def _conn(self, db):
        return connect(self.cfg, db, self.s.query_timeout_minutes)

    def count_part(self, part, by_db):
        total = 0
        for db, tables in by_db.items():
            q, params = Q.count_query(part, tables, self.types[part["dataset"]])
            conn = self._conn(db)
            try:
                with conn.cursor() as cur:
                    cur.execute(q, params)
                    total += cur.fetchone()[0]
            finally:
                conn.close()
        return total

    def _stream_db(self, part, db, tables, all_tables, with_key):
        q, params = Q.rows_query(part, tables, self.types[part["dataset"]], all_tables, with_sort_key=with_key)
        conn = self._conn(db)
        try:
            with conn.cursor(name="bulkex_rows") as cur:
                cur.itersize = 5000
                cur.execute(q, params)
                yield from cur
        finally:
            conn.close()

    def stream_rows(self, part, by_db, all_tables):
        sort = (part["output"]["rows"].get("sort") or [None])[0]
        if not sort or len(by_db) == 1:
            for db, tables in by_db.items():
                yield from self._stream_db(part, db, tables, all_tables, False)
            return
        # months spread over several databases: each comes back sorted, merge them on the sort key (blanks last)
        desc = sort["dir"] == "desc"
        if desc:
            key = lambda r: (r[-1] is not None, r[-1] if r[-1] is not None else 0)  # noqa: E731
        else:
            key = lambda r: (r[-1] is None, r[-1] if r[-1] is not None else 0)  # noqa: E731
        streams = [self._stream_db(part, db, t, all_tables, True) for db, t in by_db.items()]
        for r in heapq.merge(*streams, key=key, reverse=desc):
            yield r[:-1]

    def summary_rows(self, part, by_db, all_tables):
        types = self.types[part["dataset"]]
        top = part["output"]["summary"].get("top")
        many = len(by_db) > 1
        merged, gb, sums = {}, [], []
        for db, tables in by_db.items():
            q, params, gb, sums = Q.summary_query(part, tables, types, with_limit=not many, all_tables=all_tables)
            conn = self._conn(db)
            try:
                with conn.cursor() as cur:
                    cur.execute(q, params)
                    for r in cur.fetchall():
                        key, cnt, vals = tuple(r[:len(gb)]), r[len(gb)], list(r[len(gb) + 1:])
                        if key in merged:
                            m = merged[key]
                            m[0] += cnt
                            m[1] = [(a or 0) + (b or 0) if (a is not None or b is not None) else None for a, b in zip(m[1], vals)]
                        else:
                            merged[key] = [cnt, vals]
            finally:
                conn.close()
        rows = sorted(merged.items(), key=lambda kv: -kv[1][0])
        if top:
            rows = rows[:top]
        return gb, sums, [list(k) + [v[0]] + v[1] for k, v in rows]

    def process(self, req, status_rec):
        """Run a validated request. Returns (rows, files, message)."""
        ds_label = lambda ds: self.catalog["datasets"][ds]["label"]  # noqa: E731
        two = len(req["parts"]) > 1
        plan, notes = [], []
        for p in req["parts"]:
            ds = p["dataset"]
            by_db, missing = Q.tables_for(p, self.found[ds])
            if not by_db:
                raise Q.RequestError(f"{ds_label(ds)}: there is no data for the chosen months.")
            if missing:
                notes.append(f"{ds_label(ds)}: no data for " + ", ".join(missing))
            all_tables = [t for ts in by_db.values() for t in ts]
            plan.append((p, by_db, all_tables))

        counts = {}
        for p, by_db, _ in plan:
            counts[p["dataset"]] = self.count_part(p, by_db)
        if req["action"] == "count":
            msg = "; ".join(f"{ds_label(ds)}: {n:,} rows" for ds, n in counts.items())
            warn = []
            for p, _, _ in plan:
                n, lab = counts[p["dataset"]], ds_label(p["dataset"])
                if not p["output"].get("rows"):
                    continue
                if n > self.s.csv_max_rows:
                    warn.append(f"{lab}: over the {self.s.csv_max_rows:,} row limit - Get the file will be refused. Pick fewer months or add a filter.")
                elif n > self.s.max_rows:
                    warn.append(f"{lab}: very large - over {self.s.max_rows:,} rows, so the whole result will come as CSV files "
                                "(no Excel) and may take a long time.")
                elif n > self.s.xlsx_max_rows:
                    warn.append(f"{lab}: large - the file will need {-(-n // self.s.xlsx_max_rows)} Detail sheets and may take a while.")
            status_rec["warning"] = " ".join(warn)
            return sum(counts.values()), [], msg + (". " + "; ".join(notes) if notes else "")
        free_gb = shutil.disk_usage(self.s.workspace).free / 1024 ** 3
        if free_gb < self.s.min_free_gb:
            log.error("Only %.1f GB free on the database computer; export refused", free_gb)
            raise Q.RequestError(f"The database computer is almost out of disk space ({free_gb:.1f} GB free). "
                                 "The admin needs to free some space; then send the request again.")
        for p, _, _ in plan:
            n = counts[p["dataset"]]
            if p["output"].get("rows") and n > self.s.csv_max_rows:
                raise Q.RequestError(f"{ds_label(p['dataset'])}: {n:,} rows is over the {self.s.csv_max_rows:,} limit. "
                                     "Pick fewer months or add a filter.")

        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
        base = f"{stamp}_{slug(req.get('requester'), 20)}_{slug(req.get('name') or req['action'])}"
        tmp_dir = self.s.system_dir / "tmp"
        as_csv = any(p["output"].get("rows") and counts[p["dataset"]] > self.s.max_rows for p, _, _ in plan)
        files, total = [], 0
        if as_csv:  # too big for Excel: every part of the result goes to CSV, no Excel file at all
            for p, by_db, all_tables in plan:
                ds, out = p["dataset"], p["output"]
                tag = f"_{slug(ds_label(ds))}" if two else ""
                if out.get("rows"):
                    n = counts[ds]
                    heads = ["Year", "Month"] + [self.label(ds, c) for c in out["rows"]["columns"]]
                    files.append(_write_csv(tmp_dir / f"{base}{tag}_detail.csv", heads, self.stream_rows(p, by_db, all_tables)))
                    total += n
                if out.get("summary"):
                    gb, sums, rows = self.summary_rows(p, by_db, all_tables)
                    heads = [self.label(ds, c) for c in gb] + ["Records"] + ["Total " + self.label(ds, c) for c in sums]
                    files.append(_write_csv(tmp_dir / f"{base}{tag}_summary.csv", heads, rows))
                    if not out.get("rows"):
                        total += counts[ds]
            notes.insert(0, f"Over {self.s.max_rows:,} rows, so the result is in CSV files (no Excel)")
        else:
            wb = Workbook(write_only=True)
            info = wb.create_sheet("Request")
            _header(info, ["Item", "Value"])
            for row in [("Request", req.get("name") or ""), ("Requested by", req.get("requester") or ""),
                        ("What was asked", req.get("description") or ""), ("Made at", dt.datetime.now().strftime("%Y-%m-%d %H:%M"))]:
                info.append(_row(info, row))
            for p, by_db, all_tables in plan:
                ds = p["dataset"]
                prefix = (ds_label(ds) + " ") if two else ""
                out = p["output"]
                months = ", ".join(sorted({t["period"] for t in all_tables}))
                info.append(_row(info, (f"{ds_label(ds)} months", months)))
                if out.get("rows"):
                    heads = ["Year", "Month"] + [self.label(ds, c) for c in out["rows"]["columns"]]
                    n = counts[ds]
                    per_sheet = self.s.xlsx_max_rows  # one Excel sheet holds ~1 million rows: bigger results continue on the next sheet
                    sheets = max(1, -(-n // per_sheet))
                    name = lambda i: ((prefix + "Detail") + (f" ({i})" if sheets > 1 else ""))[:31]  # noqa: E731
                    ws, i, in_sheet = wb.create_sheet(name(1)), 1, 0
                    _header(ws, heads)
                    for r in self.stream_rows(p, by_db, all_tables):
                        if in_sheet >= per_sheet:
                            i += 1
                            ws, in_sheet = wb.create_sheet(name(i)), 0
                            _header(ws, heads)
                        ws.append(_row(ws, r))
                        in_sheet += 1
                    info.append(_row(info, (f"{prefix}Detail rows", n)))
                    if i > 1:
                        notes.append(f"{prefix}Detail is split over {i} sheets ({per_sheet:,} rows each)")
                    total += n
                if out.get("summary"):
                    gb, sums, rows = self.summary_rows(p, by_db, all_tables)
                    ws = wb.create_sheet((prefix + "Summary")[:31])
                    _header(ws, [self.label(ds, c) for c in gb] + ["Records"] + ["Total " + self.label(ds, c) for c in sums])
                    for r in rows:
                        ws.append(_row(ws, r))
                    info.append(_row(info, (f"{prefix}Summary rows", len(rows))))
                    if not out.get("rows"):
                        total += counts[ds]
            for n in notes:
                info.append(_row(info, ("Note", n)))
            xlsx = tmp_dir / f"{base}.xlsx"
            wb.save(xlsx)
            files.insert(0, xlsx)
        final = []
        for f in files:
            dest = self.s.output_dir / f.name
            i = 2
            while dest.exists():
                dest = self.s.output_dir / f"{f.stem}_{i}{f.suffix}"
                i += 1
            shutil.move(str(f), str(dest))
            free_up(dest)  # OneDrive uploads it, then frees the space on the database computer
            final.append(dest.name)
        return total, final, "; ".join(notes)

    # -- the pass --------------------------------------------------------------------------------
    def pending(self, settle_seconds=3):
        now = time.time()
        files = []
        for f in sorted(self.s.input_dir.glob("*.json"), key=lambda p: p.stat().st_mtime):
            if f.name.startswith(("~", ".")) or now - f.stat().st_mtime < settle_seconds:
                continue  # still being written / synced
            files.append(f)
        return files

    def run_once(self, settle_seconds=3):
        self.s.ensure_dirs()
        status = Status(self.s.system_dir)
        try:
            self.refresh_catalog()
        except DatabaseDown as e:
            log.error("No data tables reachable, requests stay in input/: %s", e)
            return 0
        done = 0
        files = self.pending(settle_seconds)
        if any(online_only(f) for f in files):
            keep_workspace(self.s.workspace)  # ask OneDrive to download them; they are read on a later pass
        for f in files:
            try:
                if online_only(f):
                    raise OSError("online-only in OneDrive")
                req = json.loads(f.read_text(encoding="utf-8-sig"))
            except OSError as e:  # not downloaded by OneDrive yet: wait, it is not a broken request
                if not any(r.get("id") == f.stem for r in status.active):
                    log.warning("Waiting for OneDrive to download %s (%s). If this stays, right-click the BulkEx folder "
                                "and choose 'Always keep on this device'.", f.name, e)
                status.put({"id": f.stem, "name": f.stem, "state": "waiting", "_file": f.name,
                            "message": "Waiting for OneDrive to bring this request to the database computer."})
                continue
            except ValueError as e:
                self._finish(f, status, {"id": f.stem, "name": f.stem, "state": "failed",
                                         "message": "This file is not a BulkEx request (it can't be read)."}, "failed")
                log.warning("Unreadable request %s: %s", f.name, e)
                continue
            rec = {"id": str(req.get("id") or f.stem), "name": req.get("name") or "", "requester": req.get("requester") or "",
                   "description": req.get("description") or "", "action": req.get("action"), "state": "running",
                   "rows": None, "files": [], "message": "", "started_at": dt.datetime.now().isoformat(timespec="seconds"),
                   "_file": f.name}
            status.put(rec)
            t0 = time.time()
            try:
                self.validate(req)
                rows, files, msg = self.process(req, rec)
                rec.update(state="done", rows=rows, files=files, message=msg)
                self._finish(f, status, rec, "done")
                log.info("Done %s (%s rows, %.1fs) -> %s", f.name, rows, time.time() - t0, files)
            except Q.RequestError as e:
                rec.update(state="failed", message=str(e))
                self._finish(f, status, rec, "failed")
                log.info("Refused %s: %s", f.name, e)
            except psycopg2.errors.QueryCanceled:
                rec.update(state="failed", message=f"This request took longer than {self.s.query_timeout_minutes} minutes and was stopped. "
                                                   "Pick fewer months or add a filter.")
                self._finish(f, status, rec, "failed")
                log.warning("Timeout %s", f.name)
            except psycopg2.OperationalError as e:
                rec.update(state="waiting", message="The database can't be reached right now. It will be tried again automatically.")
                status.put(rec)  # file stays in input/ and is retried next pass
                log.error("DB unreachable for %s: %s", f.name, e)
                break
            except Exception as e:  # noqa: BLE001 - one bad request must not stop the others
                rec.update(state="failed", message="Something went wrong while making this file. The admin can see the details in the log.")
                self._finish(f, status, rec, "failed")
                log.error("Error %s: %s\n%s", f.name, e, traceback.format_exc())
            done += 1
        status.save()
        return done

    def _finish(self, f, status, rec, folder):
        """Move the request to done/ or failed/ with its result inside: those two folders are the list the page shows."""
        rec["finished_at"] = dt.datetime.now().isoformat(timespec="seconds")
        rec.pop("_file", None)
        dest = self.s.system_dir / folder / f.name
        if dest.exists():
            dest = dest.with_name(f"{f.stem}_{int(time.time())}{f.suffix}")
        try:
            try:
                text = f.read_text(encoding="utf-8-sig", errors="replace")
                body = json.loads(text)
                if not isinstance(body, dict):
                    raise ValueError
            except ValueError:
                body = {"unreadable_file": text[:2000]}
            body["_result"] = rec
            dest.write_text(json.dumps(body, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
            f.unlink()
        except OSError as e:
            log.error("Could not move %s: %s", f.name, e)
        status.put(rec)
        self._log_csv(rec)

    def _log_csv(self, rec):
        path = self.s.log_dir / "requests.csv"
        new = not path.exists()
        with open(path, "a", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            if new:
                w.writerow(["finished_at", "requester", "name", "action", "state", "rows", "files", "message", "description"])
            w.writerow([rec.get("finished_at"), rec.get("requester"), rec.get("name"), rec.get("action"), rec.get("state"),
                        rec.get("rows"), ";".join(rec.get("files") or []), rec.get("message"), rec.get("description")])


def run_with_lock(runner: Runner, settle_seconds=3):
    """Only one pass at a time, even if Task Scheduler starts another while a big request is running."""
    runner.s.ensure_dirs()
    lock = runner.s.system_dir / "runner.lock"
    if lock.exists() and time.time() - lock.stat().st_mtime < LOCK_STALE_SECONDS:
        log.info("Another pass is running; skipped")
        return None
    lock.write_text(f"{os.getpid()} {dt.datetime.now().isoformat(timespec='seconds')}", encoding="utf-8")
    try:
        return runner.run_once(settle_seconds)
    finally:
        try:
            lock.unlink()
        except OSError:
            pass
