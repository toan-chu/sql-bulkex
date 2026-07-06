#!/usr/bin/env python3
"""
SQL BulkEx Portal — export dữ liệu PostgreSQL qua menu, không cần viết SQL.

Chạy:   python portal.py
Config: connection.yaml (cùng thư mục) — máy khác chỉ cần sửa file này.

Flow:  database → schema → bảng → cột → lọc → tách file → sắp xếp → review → export.
- Mọi menu đều có "← Quay lại".
- Tách file: chọn 1 cột (vd nước xuất xứ) → mỗi giá trị trong cột = 1 file riêng.
- Hàng đợi: xếp nhiều query (khác database/schema cũng được) rồi export 1 lượt.
"""

import csv
import datetime as dt
import getpass
import os
import re
import sys
import argparse

sys.dont_write_bytecode = True

import psycopg2
from psycopg2 import sql
import questionary
from questionary import Choice
import yaml

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "connection.yaml")
PASSWORD_FILE = os.path.join(BASE_DIR, ".password")
JOBS_FILE = os.path.join(BASE_DIR, "jobs.yaml")
LOG_DIR = os.path.join(BASE_DIR, "log")
PORTAL_LOG_FILE = os.path.join(LOG_DIR, "portal.log")

BACK = "<<BACK>>"
SKIP = "<<SKIP>>"  # "không chọn" — không dùng None vì trùng tín hiệu Ctrl+C của questionary
MERGE_COL = "bang_nguon"  # cột đánh dấu bảng gốc khi gộp nhiều bảng
MAX_SPLIT_FILES = 200     # tách file quá số này thì hỏi lại cho chắc
JOB_STATE_KEYS = ("db", "schema", "tables", "cols", "filters", "split", "split_len", "sort", "merged")

OPS = [
    ("bằng đúng        — khớp nguyên giá trị, nhập 1 giá trị", "eq"),
    ("trong danh sách  — nhiều giá trị, cách nhau dấu phẩy", "in"),
    ("bắt đầu bằng     — 1 hoặc nhiều prefix, cách nhau dấu phẩy (vd: BY, KZ, KG)", "prefix"),
    ("chứa chuỗi       — chuỗi nhập nằm ở bất kỳ vị trí nào", "contains"),
    ("trong khoảng     — từ X đến Y", "between"),
]


# ---------------- menu helpers ----------------

def ask(q):
    """Hỏi questionary, thoát gọn khi Ctrl+C/ESC."""
    ans = q.ask()
    if ans is None:
        print("Đã thoát.")
        sys.exit(0)
    return ans


def sel(message, choices):
    """Menu chọn 1, luôn có ← Quay lại ở cuối."""
    ch = list(choices) + [Choice(title="← Quay lại", value=BACK)]
    return ask(questionary.select(message, choices=ch))


def checkbox_back(message, choices):
    """Menu tick nhiều, dòng đầu là ← Quay lại. Trả về BACK hoặc list đã tick."""
    ch = [Choice(title="← Quay lại (tick dòng này)", value=BACK)] + list(choices)
    while True:
        picked = ask(questionary.checkbox(message, choices=ch))
        if BACK in picked:
            return BACK
        if picked:
            return picked
        print("Chưa tick gì — tick ít nhất 1 dòng (hoặc tick ← Quay lại).")


# ---------------- config & connect ----------------

def load_config():
    if not os.path.exists(CONFIG_FILE):
        print(f"Không tìm thấy {CONFIG_FILE}.")
        print("Tạo file connection.yaml (xem README) rồi chạy lại.")
        sys.exit(1)
    with open(CONFIG_FILE, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def ensure_password(cfg):
    """Load password from YAML, then .password, then prompt for portal use."""
    password = config_or_file_password(cfg)
    if not password:
        password = getpass.getpass(
            f"Password PostgreSQL (user: {cfg.get('user', 'postgres')}): "
        )
    cfg["_password"] = password
    return cfg


def password_from_file(path=None):
    path = path or PASSWORD_FILE
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8-sig") as f:
        password = f.readline().strip()
    return password or None


def config_or_file_password(cfg):
    password = str(cfg.get("password") or "").strip()
    if password:
        return password
    return password_from_file()


def connect(cfg, dbname):
    return psycopg2.connect(
        host=cfg.get("host", "localhost"),
        port=cfg.get("port", 5432),
        user=cfg.get("user", "postgres"),
        password=cfg["_password"],
        dbname=dbname,
    )


def get_conn(conns, cfg, dbname):
    """Tái sử dụng connection theo database."""
    c = conns.get(dbname)
    if c is None or c.closed:
        c = connect(cfg, dbname)
        c.set_client_encoding("UTF8")
        conns[dbname] = c
    return c


def fetch_col(cur, query, params=None):
    cur.execute(query, params)  # params=None → psycopg2 bỏ qua xử lý %, an toàn với LIKE '%'
    return [r[0] for r in cur.fetchall()]


def safe_name(s):
    """Tên file hợp lệ Windows (dbname có thể chứa '/')."""
    return re.sub(r'[\\/:*?"<>|\s]+', "_", str(s))


def append_portal_log(message):
    os.makedirs(LOG_DIR, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with open(PORTAL_LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"[{stamp}] {message}\n")


def load_saved_jobs(path=None):
    path = path or JOBS_FILE
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    return data.get("jobs", {}) or {}


def write_saved_jobs(jobs, path=None):
    path = path or JOBS_FILE
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump({"jobs": jobs}, f, allow_unicode=True, sort_keys=True)


def serializable_job_state(st):
    return {key: st.get(key) for key in JOB_STATE_KEYS}


def save_job(name, state, path=None):
    path = path or JOBS_FILE
    jobs = load_saved_jobs(path)
    jobs[name] = state
    write_saved_jobs(jobs, path)


def prompt_save_job(st):
    if not ask(questionary.confirm("Lưu job này?", default=False)):
        return
    name = ask(questionary.text("Tên job:")).strip()
    if not name:
        print("Tên job trống — bỏ qua lưu.")
        return
    save_job(name, serializable_job_state(st))
    print(f"Đã lưu job: {name}")


def list_saved_jobs():
    jobs = load_saved_jobs()
    if not jobs:
        print("Chưa có job nào trong jobs.yaml.")
        return
    print("Saved jobs:")
    for name in sorted(jobs):
        st = jobs[name]
        print(f"- {name}: {st.get('db')}.{st.get('schema')} ({len(st.get('tables') or [])} bảng)")


def missing_tables(cur, schema, tables):
    cur.execute(
        """
        SELECT table_name
        FROM information_schema.tables
        WHERE table_schema = %s
          AND table_name = ANY(%s)
        """,
        (schema, list(tables)),
    )
    existing = {r[0] for r in cur.fetchall()}
    return [t for t in tables if t not in existing]


def export_format_handler(fmt):
    return (".csv", export_csv) if str(fmt).lower().startswith("csv") else (".xlsx", export_xlsx)


# ---------------- khám phá metadata ----------------

def pick_database(cfg, conns):
    cur = get_conn(conns, cfg, cfg.get("maintenance_db", "postgres")).cursor()
    dbs = fetch_col(cur, "SELECT datname FROM pg_database WHERE datistemplate = false ORDER BY datname")
    cur.close()
    default = cfg.get("default_database")
    if default in dbs:  # đưa database mặc định lên đầu menu
        dbs = [default] + [d for d in dbs if d != default]
    return sel("Chọn database:", dbs)


def pick_schema(cur):
    schemas = fetch_col(cur, """
        SELECT schema_name FROM information_schema.schemata
        WHERE schema_name NOT LIKE 'pg\\_%' AND schema_name <> 'information_schema'
        ORDER BY schema_name""")
    if not schemas:
        print("Database này không có schema nào ngoài schema hệ thống.")
        return BACK
    return sel("Chọn schema:", schemas)


def pick_tables(cur, schema):
    tables = fetch_col(cur, """
        SELECT table_name FROM information_schema.tables
        WHERE table_schema = %s AND table_type = 'BASE TABLE'
        ORDER BY table_name""", (schema,))
    if not tables:
        print(f"Schema {schema} không có bảng nào.")
        return BACK
    return checkbox_back("Chọn bảng (Space để tick, chọn nhiều bảng sẽ gộp/tách khi export):", tables)


def get_columns(cur, schema, table):
    """[(tên cột, kiểu dữ liệu), ...] theo thứ tự trong bảng."""
    cur.execute("""
        SELECT column_name, data_type FROM information_schema.columns
        WHERE table_schema = %s AND table_name = %s
        ORDER BY ordinal_position""", (schema, table))
    return cur.fetchall()


def pick_columns(cur, schema, tables):
    """Mặc định = TẤT CẢ cột (Enter là qua). Chọn lọc bớt là ngoại lệ."""
    per_table = {t: get_columns(cur, schema, t) for t in tables}
    base = per_table[tables[0]]
    if len(tables) > 1:
        common = set.intersection(*(set(c for c, _ in cols) for cols in per_table.values()))
        base = [(c, d) for c, d in base if c in common]
        print(f"({len(base)} cột có mặt ở tất cả {len(tables)} bảng đã chọn)")
    while True:
        mode = sel("Cột export:", [
            Choice(f"Lấy TẤT CẢ {len(base)} cột (mặc định — Enter)", "all"),
            Choice("Chọn lọc bớt cột...", "some"),
        ])
        if mode == BACK:
            return BACK
        if mode == "all":
            return [name for name, _ in base]
        picked = checkbox_back("Tick cột muốn export:", [
            Choice(title=f"{name}   [{dtype}]", value=name) for name, dtype in base])
        if picked == BACK:
            continue  # về menu "Cột export"
        return picked


# ---------------- filter ----------------

def sample_values(cur, schema, table, col, n=3):
    """3 giá trị thật trong cột — cho user thấy format dữ liệu trước khi nhập."""
    try:
        cur.execute(sql.SQL(
            "SELECT DISTINCT {c}::text FROM {s}.{t} WHERE {c} IS NOT NULL LIMIT %s"
        ).format(c=sql.Identifier(col), s=sql.Identifier(schema), t=sql.Identifier(table)), (n,))
        return [r[0] for r in cur.fetchall()]
    except Exception:
        cur.connection.rollback()
        return []


def build_filters(cur, schema, tables, columns):
    """Hỏi các điều kiện lọc. Trả về BACK hoặc [(cột, op, giá_trị), ...] — nối bằng AND."""
    first = sel("Lọc dữ liệu?", [
        Choice("Không lọc — lấy hết", "no"),
        Choice("Có — thêm điều kiện lọc", "yes"),
    ])
    if first == BACK:
        return BACK
    if first == "no":
        return []
    filters = []
    while True:
        col = sel("Lọc theo cột:", columns)
        if col == BACK:
            if filters:
                print("(Hủy các điều kiện đã nhập)")
            return BACK
        samples = sample_values(cur, schema, tables[0], col)
        if samples:
            print(f"   Giá trị mẫu trong cột {col}: {', '.join(samples)}")
        op = sel("Kiểu so sánh:", [Choice(title=t, value=v) for t, v in OPS])
        if op == BACK:
            continue  # chọn lại cột
        if op in ("in", "prefix"):
            raw = ask(questionary.text("Nhập giá trị — nhiều giá trị cách nhau dấu phẩy (bỏ trống = hủy điều kiện này):"))
            val = [v.strip() for v in raw.split(",") if v.strip()]
            if not val:
                continue
        elif op == "between":
            lo = ask(questionary.text("Từ giá trị:")).strip()
            hi = ask(questionary.text("Đến giá trị:")).strip()
            if not lo or not hi:
                continue
            val = (lo, hi)
        else:
            val = ask(questionary.text("Nhập giá trị (bỏ trống = hủy điều kiện này):")).strip()
            if not val:
                continue
        filters.append((col, op, val))
        nxt = sel(f"Đã có {len(filters)} điều kiện. Tiếp?", [
            Choice("Chốt lọc — đi tiếp", "done"),
            Choice("Thêm điều kiện nữa (AND)", "more"),
        ])
        if nxt == "done":
            return filters
        if nxt == BACK:
            filters.pop()  # bỏ điều kiện vừa nhập


def where_clause(filters):
    """(fragment sql.SQL, params). ::text để so sánh an toàn với mọi kiểu cột."""
    if not filters:
        return sql.SQL(""), []
    parts, params = [], []
    for col, op, val in filters:
        ident = sql.Identifier(col)
        if op == "eq":
            parts.append(sql.SQL("{}::text = %s").format(ident))
            params.append(val)
        elif op == "prefix":
            vals = val if isinstance(val, (list, tuple)) else [val]
            ors = sql.SQL(" OR ").join(
                sql.SQL("{}::text LIKE %s").format(ident) for _ in vals)
            parts.append(sql.SQL("(") + ors + sql.SQL(")"))
            params.extend(v + "%" for v in vals)
        elif op == "contains":
            parts.append(sql.SQL("{}::text ILIKE %s").format(ident))
            params.append("%" + val + "%")
        elif op == "in":
            parts.append(sql.SQL("{}::text = ANY(%s)").format(ident))
            params.append(list(val))
        elif op == "between":
            parts.append(sql.SQL("{} BETWEEN %s AND %s").format(ident))
            params.extend(val)
    return sql.SQL(" WHERE ") + sql.SQL(" AND ").join(parts), params


# ---------------- build query ----------------

def build_query(schema, table, columns, where_sql, source_label=None):
    """SELECT cột... FROM schema.table [WHERE ...]. source_label → thêm cột bang_nguon."""
    sel_cols = sql.SQL(", ").join(sql.Identifier(c) for c in columns)
    if source_label is not None:
        sel_cols = sql.SQL("{} AS {}, ").format(
            sql.Literal(source_label), sql.Identifier(MERGE_COL)) + sel_cols
    return (sql.SQL("SELECT ") + sel_cols
            + sql.SQL(" FROM {}.{}").format(sql.Identifier(schema), sql.Identifier(table))
            + where_sql)


def build_merged_query(schema, tables, columns, where_sql):
    """UNION ALL nhiều bảng, thêm cột bang_nguon để biết dòng từ bảng nào."""
    parts = [build_query(schema, t, columns, where_sql, source_label=t) for t in tables]
    return sql.SQL(" UNION ALL ").join(parts)


def with_sort(query, sort):
    """Bọc ORDER BY ngoài cùng. sort = (cột, 'ASC'|'DESC') hoặc None."""
    if not sort:
        return query
    col, direction = sort
    return (sql.SQL("SELECT * FROM (") + query
            + sql.SQL(") _s ORDER BY {} ").format(sql.Identifier(col))
            + sql.SQL(direction))


def distinct_values(cur, schema, tables, col, filters, length=None):
    """Các giá trị khác nhau của cột (sau khi áp filter) — để tách file.
    length=N → distinct theo N ký tự đầu (vd 2 = mã nước từ mã cảng KGZZZ→KG)."""
    where_sql, params = where_clause(filters)
    if length:
        expr = sql.SQL("LEFT({}::text, {})").format(sql.Identifier(col), sql.Literal(int(length)))
    else:
        expr = sql.SQL("{}::text").format(sql.Identifier(col))
    parts = [sql.SQL("SELECT ") + expr + sql.SQL(" AS v FROM {}.{}").format(
        sql.Identifier(schema), sql.Identifier(t)) + where_sql
        for t in tables]
    q = (sql.SQL("SELECT DISTINCT v FROM (") + sql.SQL(" UNION ALL ").join(parts)
         + sql.SQL(") _d WHERE v IS NOT NULL ORDER BY v"))
    cur.execute(q, params * len(tables))
    return [r[0] for r in cur.fetchall()]


def make_jobs(st, cur):
    """Từ state đã chọn → list jobs [(tên_file, dbname, query, params), ...]."""
    db, schema, tables, cols = st["db"], st["schema"], st["tables"], st["cols"]
    merged = st["merged"]
    sort = st.get("sort")
    base = f"gop{len(tables)}bang" if merged else None

    def one(filters, suffix=""):
        where_sql, params = where_clause(filters)
        if merged:
            q = build_merged_query(schema, tables, cols, where_sql)
            return [(safe_name(f"{db}_{schema}_{base}{suffix}"), db, with_sort(q, sort), params * len(tables))]
        return [(safe_name(f"{db}_{schema}_{t}{suffix}"), db,
                 with_sort(build_query(schema, t, cols, where_sql), sort), list(params))
                for t in tables]

    if st.get("split"):
        length = st.get("split_len")
        values = distinct_values(cur, schema, tables, st["split"], st["filters"], length)
        op = "prefix" if length else "eq"  # tách theo N ký tự đầu → lọc LIKE 'XX%'
        jobs = []
        for v in values:
            jobs.extend(one(st["filters"] + [(st["split"], op, v)], suffix=f"_{v}"))
        return jobs
    return one(st["filters"])


# ---------------- preview & export ----------------

def show_preview(cur, query, params, limit=10):
    cur.execute(sql.SQL("SELECT * FROM (") + query + sql.SQL(") _p LIMIT %s"),
                list(params) + [limit])
    rows = cur.fetchall()
    headers = [d[0] for d in cur.description]
    shown = headers[:8]
    print("\n--- PREVIEW (10 dòng đầu, tối đa 8 cột đầu) ---")
    print(" | ".join(h[:18] for h in shown))
    for r in rows:
        print(" | ".join(str(v)[:18] if v is not None else "" for v in r[:8]))
    if len(headers) > 8:
        print(f"... và {len(headers) - 8} cột nữa (có đủ trong file export)")
    print("--- HẾT PREVIEW ---\n")


def count_rows(cur, query, params):
    cur.execute(sql.SQL("SELECT COUNT(*) FROM (") + query + sql.SQL(") _c"), params)
    return cur.fetchone()[0]


def export_csv(conn, query, params, headers, filepath):
    with conn.cursor(name="bulkex_stream") as cur:  # server-side cursor: không ăn hết RAM
        cur.itersize = 5000
        cur.execute(query, params)
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(headers)
            n = 0
            for row in cur:
                w.writerow(row)
                n += 1
    return n


def export_xlsx(conn, query, params, headers, filepath):
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append(headers)
    n = 0
    with conn.cursor(name="bulkex_stream") as cur:
        cur.itersize = 5000
        cur.execute(query, params)
        for row in cur:
            ws.append(list(row))
            n += 1
    wb.save(filepath)
    return n


def output_dir(cfg):
    d = cfg.get("output_dir") or os.path.join(BASE_DIR, "exports")
    os.makedirs(d, exist_ok=True)
    return d


def run_export(conns, cfg, jobs, fmt=None):
    """jobs: [(tên_file_gốc, dbname, query, params), ...] — có thể thuộc nhiều database."""
    if fmt is None:
        fmt = sel("Định dạng file:", ["CSV (.csv)", "Excel (.xlsx)"])
        if fmt == BACK:
            return False
    ext, fn = export_format_handler(fmt)
    outdir = output_dir(cfg)
    stamp = dt.date.today().strftime("%Y%m%d")
    for i, (name, dbname, query, params) in enumerate(jobs, 1):
        conn = get_conn(conns, cfg, dbname)
        filepath = os.path.join(outdir, f"{stamp}_{name}{ext}")
        with conn.cursor() as cur:
            cur.execute(sql.SQL("SELECT * FROM (") + query + sql.SQL(") _h LIMIT 0"), params)
            headers = [d[0] for d in cur.description]
        print(f"[{i}/{len(jobs)}] {os.path.basename(filepath)} ...", end=" ", flush=True)
        n = fn(conn, query, params, headers, filepath)
        conn.rollback()  # dọn transaction cho query sau
        print(f"{n:,} dòng")
    print(f"\nXong {len(jobs)} file. Nằm ở: {outdir}\n")
    return True


# ---------------- step machine ----------------

def build_job(cfg, conns, queued=0):
    """Dẫn user qua các bước, ← Quay lại lùi 1 bước.
    Trả về None (dừng ở bước database), ("drop", []) hoặc ("export"/"queue", jobs)."""
    st = {}
    step = "db"
    while True:
        if step == "db":
            r = pick_database(cfg, conns)
            if r == BACK:
                return None
            st["db"] = r
            st["cur"] = get_conn(conns, cfg, r).cursor()
            step = "schema"

        elif step == "schema":
            r = pick_schema(st["cur"])
            if r == BACK:
                st["cur"].close()
                step = "db"
                continue
            st["schema"] = r
            step = "tables"

        elif step == "tables":
            r = pick_tables(st["cur"], st["schema"])
            if r == BACK:
                step = "schema"
                continue
            st["tables"] = r
            step = "cols"

        elif step == "cols":
            r = pick_columns(st["cur"], st["schema"], st["tables"])
            if r == BACK:
                step = "tables"
                continue
            st["cols"] = r
            step = "filters"

        elif step == "filters":
            r = build_filters(st["cur"], st["schema"], st["tables"], st["cols"])
            if r == BACK:
                step = "cols"
                continue
            st["filters"] = r
            step = "split"

        elif step == "split":
            r = sel("Tách file theo giá trị 1 cột? (vd cột nước: Canada 1 file, China 1 file...)", [
                Choice("Không tách — xuất chung", SKIP),
            ] + [Choice(title=c, value=c) for c in st["cols"]])
            if r == BACK:
                step = "filters"
                continue
            st["split"] = None if r == SKIP else r
            st["split_len"] = None
            if st["split"]:
                samples = sample_values(st["cur"], st["schema"], st["tables"][0], st["split"])
                if samples:
                    print(f"   Giá trị mẫu trong cột {st['split']}: {', '.join(samples)}")
                m = sel("Tách theo:", [
                    Choice("Giá trị nguyên vẹn (vd KGZZZ, KGBIS = 2 file khác nhau)", 0),
                    Choice("2 ký tự đầu — vd mã nước từ mã cảng (KGZZZ, KGBIS → 1 file KG)", 2),
                    Choice("Số ký tự đầu tự chọn (vd 4 = nhóm HS code)...", -1),
                ])
                if m == BACK:
                    continue  # chọn lại cột tách
                if m == -1:
                    raw = ask(questionary.text("Số ký tự đầu:")).strip()
                    if not raw.isdigit() or int(raw) < 1:
                        print("Phải là số nguyên dương — chọn lại.")
                        continue
                    m = int(raw)
                st["split_len"] = m or None
            step = "sort"

        elif step == "sort":
            r = sel("Sắp xếp dòng trong file theo cột?", [
                Choice("Không sắp xếp", SKIP),
            ] + [Choice(title=c, value=c) for c in st["cols"]])
            if r == BACK:
                step = "split"
                continue
            if r == SKIP:
                st["sort"] = None
            else:
                d = sel("Chiều sắp xếp:", [
                    Choice("Tăng dần (A→Z, nhỏ→lớn)", "ASC"),
                    Choice("Giảm dần (Z→A, lớn→nhỏ)", "DESC"),
                ])
                if d == BACK:
                    continue  # hỏi lại cột sắp xếp
                st["sort"] = (r, d)
            step = "merge"

        elif step == "merge":
            if len(st["tables"]) > 1 and st["split"]:
                st["merged"] = True  # tách theo giá trị → các bảng buộc gộp (có cột bang_nguon)
                print(f"(Nhiều bảng + tách file theo giá trị → {len(st['tables'])} bảng tự gộp, có cột {MERGE_COL})")
            elif len(st["tables"]) > 1:
                r = sel(f"Đã chọn {len(st['tables'])} bảng — xuất thế nào?", [
                    Choice("Gộp 1 file (thêm cột bang_nguon)", True),
                    Choice("Mỗi bảng 1 file riêng", False),
                ])
                if r == BACK:
                    step = "sort"
                    continue
                st["merged"] = r
            else:
                st["merged"] = False
            step = "review"

        elif step == "review":
            try:
                jobs = make_jobs(st, st["cur"])
            except Exception as e:
                st["cur"].connection.rollback()
                print(f"Query lỗi: {e}")
                step = "filters"
                continue
            if not jobs:
                print("Không có giá trị nào để tách file (cột toàn NULL?). Chọn lại.")
                step = "split"
                continue
            if len(jobs) > MAX_SPLIT_FILES:
                r = sel(f"⚠ Sẽ tạo {len(jobs)} file — hơi nhiều. Tiếp tục?", [
                    Choice("Tiếp tục — tôi biết mình làm gì", "go"),
                    Choice("Chọn lại cột tách file", "redo"),
                ])
                if r != "go":
                    step = "split"
                    continue
            name, dbname, q, p = jobs[0]
            try:
                total = count_rows(st["cur"], q, p)
            except Exception as e:
                st["cur"].connection.rollback()
                print(f"Query lỗi: {e}")
                step = "filters"
                continue
            if len(jobs) > 1:
                print(f"\nSẽ xuất {len(jobs)} file. File đầu ({name}): {total:,} dòng.")
            else:
                print(f"\nQuery này: {total:,} dòng.")
            if total == 0:
                print("⚠ 0 dòng! Hay gặp: giá trị trong DB dài hơn giá trị nhập")
                print("  (vd MST lưu 13 số '0101234567890', nhập 10 số) → sửa lọc, chọn 'bắt đầu bằng'.")
            while True:
                act = sel("Tiếp theo?", [
                    Choice(f"Export ngay ({queued + len(jobs)} file)", "export"),
                    Choice("Để vào hàng đợi — làm thêm query khác", "queue"),
                    Choice("Xem preview 10 dòng (file đầu)", "prev"),
                    Choice("Sửa lọc", "fix"),
                    Choice("Bỏ query này", "drop"),
                ])
                if act == "prev":
                    try:
                        show_preview(st["cur"], q, p)
                    except Exception as e:
                        st["cur"].connection.rollback()
                        print(f"Preview lỗi: {e}")
                    continue
                break
            if act == "fix":
                step = "filters"
                continue
            if act == BACK:
                step = "sort"
                continue
            st["cur"].close()
            if act == "drop":
                return ("drop", [])
            if act in ("export", "queue"):
                prompt_save_job(st)
            return (act, jobs)


# ---------------- main ----------------

def run_saved_job(cfg, conns, name):
    saved = load_saved_jobs()
    st = saved.get(name)
    if not st:
        print(f"Không tìm thấy job: {name}")
        return 1

    db = st["db"]
    conn = get_conn(conns, cfg, db)
    cur = conn.cursor()
    try:
        missing = missing_tables(cur, st["schema"], st["tables"])
        if missing:
            msg = f"Job {name} skipped: missing tables {', '.join(missing)}"
            append_portal_log(msg)
            print(msg)
            return 1
        st = dict(st)
        st["cur"] = cur
        jobs = make_jobs(st, cur)
    except Exception as e:
        conn.rollback()
        msg = f"Job {name} skipped: {e}"
        append_portal_log(msg)
        print(msg)
        return 1
    finally:
        cur.close()

    if not jobs:
        msg = f"Job {name} skipped: no export jobs built"
        append_portal_log(msg)
        print(msg)
        return 1

    fmt = cfg.get("job_export_format", "xlsx")
    return 0 if run_export(conns, cfg, jobs, fmt=fmt) else 1


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="SQL BulkEx Portal")
    parser.add_argument("--job", help="Chạy saved job trong jobs.yaml không cần menu")
    parser.add_argument("--list-jobs", action="store_true", help="Liệt kê saved jobs")
    return parser.parse_args(argv)


def configure_stdio():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def main(argv=None):
    configure_stdio()
    args = parse_args(argv)
    print("=== SQL BulkEx Portal ===")
    if args.list_jobs:
        list_saved_jobs()
        return 0

    cfg = ensure_password(load_config())
    conns = {}   # dbname → connection (tái sử dụng)
    queue = []   # hàng đợi jobs chờ export
    try:
        if args.job:
            return run_saved_job(cfg, conns, args.job)

        while True:
            r = build_job(cfg, conns, queued=len(queue))

            if r is None:  # ← Quay lại ở bước chọn database = muốn dừng
                if queue:
                    act = sel(f"Đang có {len(queue)} file chờ trong hàng đợi:", [
                        Choice("Export hàng đợi rồi thoát", "export"),
                        Choice("Thoát, bỏ hàng đợi", "quit"),
                    ])
                    if act == "export":
                        run_export(conns, cfg, queue)
                    elif act == BACK:
                        continue
                break

            action, jobs = r
            if action == "drop":
                continue
            queue.extend(jobs)
            if action == "export":
                if run_export(conns, cfg, queue):
                    queue = []
                if not ask(questionary.confirm("Query tiếp?", default=False)):
                    break
    finally:
        for c in conns.values():
            try:
                c.close()
            except Exception:
                pass
    print("Bye.")


if __name__ == "__main__":
    try:
        code = main()
        if isinstance(code, int):
            sys.exit(code)
    except KeyboardInterrupt:
        print("\nĐã thoát.")
    except psycopg2.OperationalError as e:
        print(f"\nLỗi kết nối database: {e}")
        print("Kiểm tra connection.yaml (host/port/user/dbname) và password.")
