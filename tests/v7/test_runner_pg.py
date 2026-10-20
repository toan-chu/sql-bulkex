"""BulkEx v7 runner against a real PostgreSQL.

Needs BULKEX_TEST_PG, e.g.  host=localhost user=postgres password=secret port=5432
Without it every test here is skipped (the v6 tests still run).
"""
import json
import os
import sys
import time
from decimal import Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "app"))

psycopg2 = pytest.importorskip("psycopg2")
from openpyxl import load_workbook  # noqa: E402

from bulkex import query as Q  # noqa: E402
from bulkex.config import Settings, connect  # noqa: E402
from bulkex.runner import Runner, run_with_lock  # noqa: E402
from bulkex.scan import scan, table_period  # noqa: E402

PG = os.environ.get("BULKEX_TEST_PG")
pytestmark = pytest.mark.skipif(not PG, reason="set BULKEX_TEST_PG to run the PostgreSQL tests")
DBS = ("bx_test_main", "bx_test_old")


def pg_cfg():
    cfg = dict(kv.split("=", 1) for kv in PG.split())
    cfg["_password"] = cfg.pop("password", "")
    cfg["maintenance_db"] = "postgres"
    return cfg


EXPORT_COLS = """so_to_khai text, ma_so_hang_hoa text, mo_ta_hang_hoa text, ma_nguoi_xuat_khau text,
    ten_nguoi_xuat_khau text, ten_nguoi_nhap_khau text, ma_nuoc text, tri_gia_hoa_don numeric, so_luong_1 numeric"""
IMPORT_COLS = """so_to_khai text, ma_so_hang_hoa text, mo_ta_hang_hoa text, ma_nguoi_nhap_khau text,
    ten_nguoi_nhap_khau text, ten_nguoi_xuat_khau text, ma_nuoc_xuat_xu text, tri_gia_hoa_don text"""


@pytest.fixture(scope="module")
def pg():
    cfg = pg_cfg()
    admin = psycopg2.connect(host=cfg["host"], port=cfg.get("port", 5432), user=cfg["user"], password=cfg["_password"], dbname="postgres")
    admin.autocommit = True
    with admin.cursor() as cur:
        for db in DBS:
            cur.execute(f"DROP DATABASE IF EXISTS {db}")
            cur.execute(f"CREATE DATABASE {db}")
    def run(db, statements):
        c = psycopg2.connect(host=cfg["host"], port=cfg.get("port", 5432), user=cfg["user"], password=cfg["_password"], dbname=db)
        c.autocommit = True
        with c.cursor() as cur:
            for s, *p in statements:
                cur.execute(s, *(p or []))
        c.close()
    rows_jan = [
        ("TK1", "84713020", "Laptop Dell 14in", "0101234567", "CTY A", "DELL INC", "US", 1500, 10),
        ("TK2", "84713090", "Notebook computer", "0101234567", "CTY A", "ACME", "CN", 900, 5),
        ("TK3", "85171200", "Mobile phone", "0309999999", "CTY B", "ACME", None, 300, 100),
        ("TK4", "84733010", "=SUM(A1:A9) parts", "0309999999", "CTY B", "PARTS LTD", "JP", 50, 1000),
        ("TK5", "39239090", "Bag 50% off", "0101234567", "CTY A", "SHOP", "US", 20, 7),
        ("TK6", "39239090", "Bag 500 off", "0101234567", "CTY A", "SHOP", "US", 20, 7),
        ("TK8", "30847100", "Medicine", "0309999999", "CTY B", "PHARMA", "DE", 10, 1),   # has 8471 inside, must NOT match "starts with 8471"
    ]
    rows_feb = [("TK7", "84713020", "Laptop HP", "0101234567", "CTY A", "HP", "SG", 1200, 3)]
    run("bx_test_main", [
        ("CREATE SCHEMA vietnam_export",), ("CREATE SCHEMA vietnam_import",), ("CREATE SCHEMA other",),
        (f"CREATE TABLE vietnam_export.x_y2025_01 ({EXPORT_COLS})",),
        (f"CREATE TABLE vietnam_export.x_y2025_02 ({EXPORT_COLS})",),
        (f"CREATE TABLE vietnam_export.x_y2025_02_backup ({EXPORT_COLS})",),   # same month twice -> warning, shortest name wins
        (f"CREATE TABLE vietnam_export.x_latest ({EXPORT_COLS})",),            # no month in the name -> warning, skipped
        ("CREATE TABLE other.customers (id int, name text)",),                  # not a data set -> ignored
        (f"CREATE TABLE vietnam_import.i_2025_1 ({IMPORT_COLS})",),
        *[("INSERT INTO vietnam_export.x_y2025_01 VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)", r) for r in rows_jan],
        *[("INSERT INTO vietnam_export.x_y2025_02 VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)", r) for r in rows_feb],
        ("INSERT INTO vietnam_export.x_y2025_02_backup VALUES ('X','1','x','1','x','x','x',1,1)",),
        ("INSERT INTO vietnam_import.i_2025_1 VALUES ('IM1','84713020','Laptop','0101234567','CTY A','LENOVO','CN','1,234.50')",),
        ("INSERT INTO vietnam_import.i_2025_1 VALUES ('IM2','85171200','Phone','0202020202','CTY C','SAMSUNG','KR','99')",),
        ("INSERT INTO vietnam_import.i_2025_1 VALUES ('IM3','85171200','Phone','0202020202','CTY C','SAMSUNG','KR','')",),
    ])
    run("bx_test_old", [  # another database with a different schema name: still found
        ("CREATE SCHEMA xk",), (f"CREATE TABLE xk.exp202412 ({EXPORT_COLS})",),
        ("INSERT INTO xk.exp202412 VALUES ('OLD1','84713020','Old laptop','0101234567','CTY A','DELL INC','US',700,2)",),
    ])
    yield cfg
    with admin.cursor() as cur:
        for db in DBS:
            cur.execute(f"DROP DATABASE IF EXISTS {db} WITH (FORCE)")
    admin.close()


@pytest.fixture
def runner(pg, tmp_path):
    s = Settings(workspace=tmp_path / "ws", log_dir=tmp_path / "log", databases=list(DBS))
    cfgdir = tmp_path / "config"
    cfgdir.mkdir()
    return Runner(s, pg, cfgdir)


def cond(field, op, *values):
    return {"field": field, "op": op, "values": list(values)}


def request(parts, action="export", name="t", rid=None):
    return {"schema": "bulkex-request/1", "id": rid or f"r{time.time_ns()}", "requester": "Tester", "name": name,
            "action": action, "description": "test", "parts": parts}


def part(ds, months, items=(), match="all", columns=("so_to_khai",), summary=None, rows=True, sort=None):
    return {"dataset": ds, "period": {"months": list(months)}, "filter": {"match": match, "items": list(items)},
            "output": {"rows": {"columns": list(columns), "sort": sort or [], "distinct": False} if rows else None,
                       "summary": summary}}


def submit(runner, req):
    runner.s.ensure_dirs()
    (runner.s.input_dir / f"{req['id']}.json").write_text(json.dumps(req), encoding="utf-8")
    runner.run_once(settle_seconds=0)
    from bulkex.runner import Status
    return next(r for r in Status(runner.s.system_dir).finished() if r["id"] == req["id"])


def sheet(runner, rec, name, keep_period=False):
    wb = load_workbook(runner.s.output_dir / rec["files"][0], read_only=True)
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    if "Detail" in name and not keep_period:
        assert rows[0][:2] == ["Year", "Month"]                        # every Detail sheet starts with Year, Month
        rows = [r[2:] for r in rows]
    return rows


# ------------------------------------------------------------------ scanning
def test_period_from_table_names():
    assert table_period("x_y2026_06") == "2026-06"
    assert table_period("i_2025_7") == "2025-07"
    assert table_period("exp202412") == "2024-12"
    assert table_period("x_y2026_06_backup") == "2026-06"
    assert table_period("x_latest") is None
    assert table_period("t_2026_13") is None


def test_scan_finds_tables_by_columns_in_any_database(pg):
    defs = {"export": {"signature": ["so_to_khai", "ma_so_hang_hoa", "ma_nguoi_xuat_khau"]},
            "import": {"signature": ["so_to_khai", "ma_so_hang_hoa", "ma_nguoi_nhap_khau"]}}
    found, warnings = scan(pg, defs, list(DBS))
    assert [(t["db"], t["table"]) for t in found["export"]] == [
        ("bx_test_old", "exp202412"), ("bx_test_main", "x_y2025_01"), ("bx_test_main", "x_y2025_02")]
    assert [t["table"] for t in found["import"]] == ["i_2025_1"]
    text = " | ".join(warnings)
    assert "x_latest" in text and "no year and month" in text
    assert "x_y2025_02_backup" in text


def test_catalog_lists_months_and_hides_database_names(runner):
    runner.s.ensure_dirs()
    runner.refresh_catalog()
    js = (runner.s.system_dir / "catalog.js").read_text(encoding="utf-8")
    cat = json.loads(js.split("= ", 1)[1].rstrip().rstrip(";"))
    assert cat["datasets"]["export"]["periods"] == ["2024-12", "2025-01", "2025-02"]
    assert "bx_test_main" not in js and "vietnam_export" not in js
    labels = {f["key"]: f["label"] for f in cat["datasets"]["export"]["fields"]}
    assert labels["ma_so_hang_hoa"] == "HS Code"
    types = {f["key"]: f["type"] for f in cat["datasets"]["import"]["fields"]}
    assert types["tri_gia_hoa_don"] == "number"   # text column listed in numbers_stored_as_text


# ------------------------------------------------------------------ filters -> rows
def test_hs_prefixes_of_any_length_and_filtered_fields_in_file(runner):
    req = request([part("export", ["2025-01"], [cond("ma_so_hang_hoa", "starts_with", "8471", "847330")],
                        columns=["so_to_khai"])])
    rec = submit(runner, req)
    assert rec["state"] == "done", rec
    rows = sheet(runner, rec, "Detail")
    assert rows[0] == ["Declaration Number", "HS Code"]           # the filtered field was added to the file
    assert sorted(r[0] for r in rows[1:]) == ["TK1", "TK2", "TK4"]


def test_is_not_keeps_blank_countries_and_groups_combine(runner):
    flt = [cond("ma_nuoc", "is_none_of", "cn"),                   # case-insensitive, keeps NULL country (TK3)
           {"match": "any", "items": [cond("ma_so_hang_hoa", "starts_with", "8517"), cond("ten_nguoi_nhap_khau", "contains", "dell")]}]
    rec = submit(runner, request([part("export", ["2025-01"], flt)]))
    assert sorted(r[0] for r in sheet(runner, rec, "Detail")[1:]) == ["TK1", "TK3"]


def test_percent_sign_is_literal_and_formula_stays_text(runner):
    rec = submit(runner, request([part("export", ["2025-01"], [cond("mo_ta_hang_hoa", "contains", "50%")])]))
    assert [r[0] for r in sheet(runner, rec, "Detail")[1:]] == ["TK5"]
    rec = submit(runner, request([part("export", ["2025-01"], [cond("so_to_khai", "is_any_of", "TK4")], columns=["mo_ta_hang_hoa"])]))
    assert sheet(runner, rec, "Detail")[1][0] == "=SUM(A1:A9) parts"


def test_numbers_stored_as_text_can_be_compared(runner):
    rec = submit(runner, request([part("import", ["2025-01"], [cond("tri_gia_hoa_don", "between", "1000", "2000")])]))
    assert [r[0] for r in sheet(runner, rec, "Detail")[1:]] == ["IM1"]
    rec = submit(runner, request([part("import", ["2025-01"], [cond("tri_gia_hoa_don", "is_empty")])]))
    assert [r[0] for r in sheet(runner, rec, "Detail")[1:]] == ["IM3"]


def test_months_across_databases_with_sort(runner):
    req = request([part("export", ["2024-12", "2025-02"], [cond("ma_so_hang_hoa", "starts_with", "8471")],
                        columns=["so_to_khai", "tri_gia_hoa_don"], sort=[{"field": "tri_gia_hoa_don", "dir": "desc"}])])
    rec = submit(runner, req)
    assert rec["state"] == "done", rec
    assert [r[0] for r in sheet(runner, rec, "Detail")[1:]] == ["TK7", "OLD1"]


# ------------------------------------------------------------------ both tables, summary, count
def test_export_and_import_parts_with_their_own_filters_and_summary(runner):
    summ = {"group_by": ["ten_nguoi_nhap_khau"], "measures": [{"fn": "count"}, {"fn": "sum", "field": "tri_gia_hoa_don"}], "top": 2}
    req = request([
        part("export", ["2025-01"], [cond("ma_nguoi_xuat_khau", "starts_with", "0101")], summary=summ),
        part("import", ["2025-01"], [cond("ma_nguoi_nhap_khau", "starts_with", "0101")]),
    ])
    rec = submit(runner, req)
    assert rec["state"] == "done", rec
    wb = load_workbook(runner.s.output_dir / rec["files"][0], read_only=True)
    assert wb.sheetnames == ["Request", "Export Detail", "Export Summary", "Import Detail"]
    wb.close()
    summary = sheet(runner, rec, "Export Summary")
    assert summary[0] == ["Importer Name", "Records", "Total Invoice Value"]
    assert summary[1] == ["SHOP", 2, 40] and len(summary) == 3          # top 2 groups, most records first
    assert [r[0] for r in sheet(runner, rec, "Import Detail")[1:]] == ["IM1"]


def test_count_only_makes_no_file(runner):
    rec = submit(runner, request([part("export", ["2025-01", "2025-02"], [cond("ma_so_hang_hoa", "starts_with", "8471")])], action="count"))
    assert rec["state"] == "done" and rec["files"] == [] and rec["rows"] == 3
    assert "3 rows" in rec["message"]


def test_missing_month_is_noted_not_fatal(runner):
    rec = submit(runner, request([part("export", ["2025-01", "2023-05"], [cond("so_to_khai", "is_any_of", "TK1")])]))
    assert rec["state"] == "done" and "2023-05" in rec["message"]


# ------------------------------------------------------------------ refusals and limits
@pytest.mark.parametrize("bad, words", [
    (cond('so_to_khai"; DROP TABLE vietnam_export.x_y2025_01; --', "is_any_of", "x"), "does not exist"),
    (cond("tri_gia_hoa_don", "contains", "1"), "can't be used"),
    (cond("tri_gia_hoa_don", "gt", "abc"), "not a number"),
    (cond("ma_so_hang_hoa", "starts_with"), "at least one value"),
])
def test_bad_requests_are_refused_with_a_reason(runner, bad, words):
    rec = submit(runner, request([part("export", ["2025-01"], [bad])]))
    assert rec["state"] == "failed" and words in rec["message"]
    assert (runner.s.system_dir / "failed" / f"{rec['id']}.json").exists()


def test_quote_in_a_value_is_just_text(runner, pg):
    rec = submit(runner, request([part("export", ["2025-01"], [cond("so_to_khai", "is_any_of", "x'); DROP TABLE vietnam_export.x_y2025_01; --")])]))
    assert rec["state"] == "done" and rec["rows"] == 0
    conn = connect(pg, "bx_test_main")
    with conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM vietnam_export.x_y2025_01")
        assert cur.fetchone()[0] == 7
    conn.close()


def test_row_limit_refuses_before_making_the_file(runner):
    runner.s.max_rows, runner.s.csv_max_rows = 1, 2
    rec = submit(runner, request([part("export", ["2025-01"])]))
    assert rec["state"] == "failed" and "over the 2 limit" in rec["message"]
    assert not list(runner.s.output_dir.iterdir())


def test_results_too_big_for_excel_come_as_csv_only(runner):
    runner.s.max_rows = 5                                              # export Jan has 7 rows
    summ = {"group_by": ["ma_nuoc"], "measures": [], "top": None}
    rec = submit(runner, request([part("export", ["2025-01"], columns=("so_to_khai", "ma_so_hang_hoa"), summary=summ),
                                  part("import", ["2025-01"])]))
    assert rec["state"] == "done" and "CSV files (no Excel)" in rec["message"]
    assert not [f for f in rec["files"] if f.endswith(".xlsx")]
    assert sorted(f.split("_t_")[-1] for f in rec["files"]) == ["Export_detail.csv", "Export_summary.csv", "Import_detail.csv"]
    detail = next(f for f in rec["files"] if f.endswith("Export_detail.csv"))
    raw = (runner.s.output_dir / detail).read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")                           # BOM: Excel reads Vietnamese right
    lines = raw.decode("utf-8-sig").splitlines()
    assert len(lines) == 8                                             # header + 7 rows


def test_more_rows_than_a_sheet_continue_on_next_sheets(runner):
    runner.s.xlsx_max_rows = 3
    rec = submit(runner, request([part("export", ["2025-01"])]))       # 7 rows -> 3 + 3 + 1
    assert rec["state"] == "done" and len(rec["files"]) == 1 and "split over 3 sheets" in rec["message"]
    wb = load_workbook(runner.s.output_dir / rec["files"][0], read_only=True)
    assert wb.sheetnames == ["Request", "Detail (1)", "Detail (2)", "Detail (3)"]
    sizes = [len(list(wb[s].iter_rows())) - 1 for s in wb.sheetnames[1:]]
    wb.close()
    assert sizes == [3, 3, 1]


def test_count_warns_about_big_results_instead_of_failing(runner):
    runner.s.xlsx_max_rows, runner.s.max_rows = 3, 5
    rec = submit(runner, request([part("export", ["2025-01"])], action="count"))
    runner.s.csv_max_rows = 6
    assert "Get the file will be refused" in submit(runner, request([part("export", ["2025-01"])], action="count"))["warning"]
    runner.s.csv_max_rows = 100
    rec = submit(runner, request([part("export", ["2025-01"])], action="count"))
    assert rec["state"] == "done" and rec["rows"] == 7 and "CSV files" in rec["warning"]
    runner.s.max_rows = 100
    rec = submit(runner, request([part("export", ["2025-01"])], action="count"))
    assert "3 Detail sheets" in rec["warning"]


@pytest.mark.parametrize("field", ["name", "requester"])
def test_request_without_names_is_refused(runner, field):
    req = request([part("export", ["2025-01"])])
    req[field] = "  "
    rec = submit(runner, req)
    assert rec["state"] == "failed" and "request name and your name" in rec["message"]


def test_connection_is_read_only_with_a_time_limit(pg):
    conn = connect(pg, "bx_test_main", timeout_minutes=7)
    with conn.cursor() as cur:
        cur.execute("SHOW statement_timeout")
        assert cur.fetchone()[0] == "7min"
        with pytest.raises(psycopg2.errors.ReadOnlySqlTransaction):
            cur.execute("CREATE TABLE vietnam_export.hack (x int)")
    conn.close()


def test_second_pass_waits_while_one_is_running(runner):
    runner.s.ensure_dirs()
    (runner.s.system_dir / "runner.lock").write_text("123", encoding="utf-8")
    assert run_with_lock(runner, settle_seconds=0) is None
    (runner.s.system_dir / "runner.lock").unlink()
    assert run_with_lock(runner, settle_seconds=0) == 0


def test_validator_converts_numbers():
    p = part("export", ["2025-01"], [cond("tri_gia_hoa_don", "between", "1,000", "5")])
    Q.validate_part(p, {"types": {"tri_gia_hoa_don": "number", "so_to_khai": "text"}}, {})
    assert p["filter"]["items"][0]["values"] == [Decimal("1000"), Decimal("5")]


def test_recent_requests_follow_the_done_and_failed_folders(runner):
    ok = submit(runner, request([part("export", ["2025-01"])], action="count", rid="ok1"))
    bad = submit(runner, request([part("export", ["1999-01"])], rid="bad1"))
    assert (ok["state"], bad["state"]) == ("done", "failed")
    def shown():
        js = (runner.s.system_dir / "status.js").read_text(encoding="utf-8")
        return [r["id"] for r in json.loads(js[js.index("{"):js.rindex("}") + 1])["requests"]]
    assert shown() == ["bad1", "ok1"]
    (runner.s.system_dir / "failed" / "bad1.json").unlink()            # the admin deletes a file ...
    runner.run_once(settle_seconds=0)
    assert shown() == ["ok1"]                                          # ... and nobody sees it any more
    for f in (runner.s.system_dir / "done").iterdir():
        f.unlink()
    runner.run_once(settle_seconds=0)
    assert shown() == []


def test_waiting_request_shows_until_its_file_is_gone(tmp_path):
    from bulkex.runner import Status
    sysdir = tmp_path / "_system"
    for d in ("done", "failed"):
        (sysdir / d).mkdir(parents=True)
    (tmp_path / "input").mkdir()
    (tmp_path / "input" / "w.json").write_text("{}", encoding="utf-8")
    st = Status(sysdir)
    st.put({"id": "w", "state": "waiting", "_file": "w.json"})
    assert [r["id"] for r in Status(sysdir).active] == ["w"]
    (tmp_path / "input" / "w.json").unlink()
    st = Status(sysdir)
    st.save()
    assert st.items == []


def test_logs_go_to_the_shared_folder_not_the_repo(tmp_path):
    s = Settings(workspace=tmp_path / "ws")
    s.ensure_dirs()
    assert s.log_dir == tmp_path / "ws" / "_system" / "log" and s.log_dir.is_dir()


def test_database_down_keeps_field_list_and_requests(runner):
    runner.s.ensure_dirs()
    runner.refresh_catalog()
    before = (runner.s.system_dir / "catalog.js").read_text(encoding="utf-8")
    runner.s.databases = ["bx_test_does_not_exist"]
    req = request([part("export", ["2025-01"])])
    (runner.s.input_dir / "r.json").write_text(json.dumps(req), encoding="utf-8")
    assert runner.run_once(settle_seconds=0) == 0
    assert (runner.s.system_dir / "catalog.js").read_text(encoding="utf-8") == before
    assert (runner.s.input_dir / "r.json").exists()          # retried on the next pass, not failed


def test_online_only_request_waits_instead_of_failing(runner, monkeypatch):
    import bulkex.runner as R
    calls = []
    monkeypatch.setattr(R, "online_only", lambda f: True)              # OneDrive has not downloaded it yet
    monkeypatch.setattr(R, "keep_workspace", lambda folder: calls.append(folder) or True)
    req = request([part("export", ["2025-01"])], rid="cloud1")
    runner.s.ensure_dirs()
    (runner.s.input_dir / "cloud1.json").write_text(json.dumps(req), encoding="utf-8")
    runner.run_once(settle_seconds=0)
    runner.run_once(settle_seconds=0)
    assert (runner.s.input_dir / "cloud1.json").exists()               # kept, retried later
    assert not list((runner.s.system_dir / "failed").iterdir())
    assert calls == [runner.s.workspace, runner.s.workspace]           # asked OneDrive to keep the folder on the computer
    js = (runner.s.system_dir / "status.js").read_text(encoding="utf-8")
    shown = json.loads(js[js.index("{"):js.rindex("}") + 1])["requests"]
    assert [(r["id"], r["state"]) for r in shown] == [("cloud1", "waiting")]
    monkeypatch.setattr(R, "online_only", lambda f: False)             # downloaded now
    runner.run_once(settle_seconds=0)
    assert (runner.s.system_dir / "done" / "cloud1.json").exists()


def test_same_scan_notes_are_logged_once(runner, caplog):
    runner.s.ensure_dirs()
    with caplog.at_level("WARNING", logger="bulkex"):
        runner.refresh_catalog()
        first = len(caplog.records)
        runner.refresh_catalog()
    assert first > 0 and len(caplog.records) == first


def test_keep_workspace_pins_requests_and_frees_results(tmp_path):
    from bulkex.onedrive import keep_workspace
    (tmp_path / "output").mkdir()
    seen = []
    ok = keep_workspace(tmp_path, run=lambda a, **k: seen.append(a[1:]) or type("R", (), {"returncode": 0})(), windows=True)
    assert ok
    assert ["-P", str(tmp_path)] in seen                                 # the folder itself is not pinned
    assert ["+P", "-U", str(tmp_path / "input" / "*"), "/S", "/D"] in seen
    assert ["-P", "+U", str(tmp_path / "output" / "*"), "/S", "/D"] in seen
    assert not [a for a in seen if a[:2] == ["+P", "-U"] and "output" in a[2]]


def test_export_refused_when_disk_is_almost_full(runner, monkeypatch):
    import bulkex.runner as R
    monkeypatch.setattr(R.shutil, "disk_usage", lambda p: type("U", (), {"free": 2 * 1024 ** 3})())
    rec = submit(runner, request([part("export", ["2025-01"])]))
    assert rec["state"] == "failed" and "out of disk space (2.0 GB free)" in rec["message"]
    assert submit(runner, request([part("export", ["2025-01"])], action="count"))["state"] == "done"   # counting needs no disk


def test_every_detail_row_says_its_year_and_month(runner):
    rec = submit(runner, request([part("export", ["2024-12", "2025-01"], columns=("so_to_khai",),
                                       sort=[{"field": "so_to_khai", "dir": "asc"}])]))
    rows = sheet(runner, rec, "Detail", keep_period=True)
    assert rows[0] == ["Year", "Month", "Declaration Number"]
    by_tk = {r[2]: (r[0], r[1]) for r in rows[1:]}
    assert by_tk["OLD1"] == (2024, 12) and by_tk["TK1"] == (2025, 1)    # from two databases, numbers not text


def test_same_column_can_have_a_different_name_per_data_set():
    from bulkex.catalog import build_catalog
    tbl = {"db": "d", "schema": "s", "table": "t", "period": "2025-01", "columns": {"ma_nuoc": "text"}}
    labels = {"fields": {"ma_nuoc": {"en": {"export": "destination_country_code", "import": "shipped_from_country_code"}}}}
    defs = {"datasets": {"export": {"label": "Export"}, "import": {"label": "Import"}}}
    cat = build_catalog({"export": [tbl], "import": [tbl]}, labels, defs)
    names = {ds: cat["datasets"][ds]["fields"][0]["label"] for ds in ("export", "import")}
    assert names == {"export": "Destination Country Code", "import": "Shipped From Country Code"}
